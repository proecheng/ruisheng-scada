"""Build the user-requested Word report from retained September 7 evidence.

Preset: compact_reference_guide; header: memo_masthead (no border).
Named overrides: Chinese fonts (Microsoft YaHei East Asian + Calibri Latin),
landscape_point_appendix (Letter landscape, 12960 DXA), table_text (9.5 pt,
1.10 line, 0/0 spacing), metadata (9.5 pt, 0/3, 1.15), code_text (9 pt,
1.15), title (23 pt, 0/5), note (9.5 pt, 4/4), footer (9 pt, 0/0).
Chinese_line_metrics overrides font-dependent multiple leading: body 17 pt,
table 13 pt, metadata 13.5 pt, notes 13 pt, title 29 pt, subtitle 16 pt,
headings 21/18/17 pt, code/footer/header 12 pt, action lists 17 pt.
"""

from __future__ import annotations

import hashlib
import json
import sys
from datetime import UTC, datetime
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path("D:/江苏润盛")
EVIDENCE = ROOT / "docs/superpowers/specs/evidence/point-release-20260907"
PREVIOUS = EVIDENCE.parent / "point-pipeline-20260907"
QA = Path(__file__).parent
OUTPUT = ROOT / "docs/reports/江苏润盛_点位表与远程测试报告_20260907.docx"
SKILL = Path(
    "C:/Users/admin/.codex/plugins/cache/openai-primary-runtime/documents/26.826.12353/skills/documents"
)
sys.path.insert(0, str(SKILL / "scripts"))
from table_geometry import apply_table_geometry, audit_docx_tables


def read_json(path):
    return json.loads(path.read_text(encoding="utf-8-sig"))


result = read_json(EVIDENCE / "target-pipeline-results.json")
report = result["reports"][0]
state = read_json(EVIDENCE / "target-final-state.json")
release = read_json(EVIDENCE / "release-build.json")
validation = read_json(EVIDENCE / "validation-summary.json")
candidates = read_json(PREVIOUS / "candidate-summary.json")
audit = [
    json.loads(line)
    for line in (EVIDENCE / "point-release-20260907-physical.jsonl")
    .read_text(encoding="utf-8")
    .splitlines()
]
raw = {p["address"]: p for p in report["physical_replay"]}
tested = {p["id"]: p for p in report["points"]}
assert len(candidates) == 46 and len(raw) == 36 and len(tested) == 46
assert validation["passed"] == 57 and validation["failed"] == 0
assert result["failure"] is None and report["mode"] == "deployed"
assert report["database_counts"] == {"realtime": 126, "history": 168}

doc = Document()
doc.core_properties.title = "江苏润盛：点位表与远程测试报告"
doc.core_properties.subject = "2026年9月7日签名部署、只读采集、数据库验证与未决问题"
doc.core_properties.author = "项目测试记录"
doc.core_properties.last_modified_by = "项目测试记录"
doc.core_properties.created = datetime(2026, 9, 7, tzinfo=UTC)
doc.core_properties.modified = datetime(2026, 9, 7, tzinfo=UTC)
doc.core_properties.keywords = "候选点位表,只读采集,独立测试库,未完成生产验收"
settings = doc.settings.element
update = OxmlElement("w:updateFields")
update.set(qn("w:val"), "true")
settings.append(update)


def style(name, size, before=0, after=6, line=1.25, color="202020", bold=False):
    s = (
        doc.styles[name]
        if name in doc.styles
        else doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    )
    s.font.name = "Calibri"
    s.font.size = Pt(size)
    s.font.bold = bold
    s.font.italic = False
    s.font.underline = False
    s.font.color.rgb = RGBColor.from_string(color)
    fonts = s.element.get_or_add_rPr().get_or_add_rFonts()
    for key in ("ascii", "hAnsi", "cs"):
        fonts.set(qn("w:" + key), "Calibri")
    fonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    for key in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        fonts.attrib.pop(qn("w:" + key), None)
    lang = OxmlElement("w:lang")
    lang.set(qn("w:val"), "en-US")
    lang.set(qn("w:eastAsia"), "zh-CN")
    s.element.get_or_add_rPr().append(lang)
    fmt = s.paragraph_format
    fmt.space_before, fmt.space_after = Pt(before), Pt(after)
    leading = {
        "Normal": 17,
        "Title": 29,
        "Subtitle": 16,
        "Heading 1": 21,
        "Heading 2": 18,
        "Heading 3": 17,
        "Table Text": 13,
        "Table Header": 13,
        "Metadata": 13.5,
        "Note": 13,
        "Table Source": 13,
        "Code Text": 12,
        "Header": 12,
        "Footer": 12,
        "Action List": 17,
        "Lead": 17,
    }
    fmt.line_spacing = Pt(leading[name])
    fmt.widow_control = True
    fmt.alignment = WD_ALIGN_PARAGRAPH.LEFT
    ppr = s.element.get_or_add_pPr()
    for border in list(ppr.findall(qn("w:pBdr"))):
        ppr.remove(border)
    snap = OxmlElement("w:snapToGrid")
    snap.set(qn("w:val"), "0")
    ppr.append(snap)
    return s


style("Normal", 11)
style("Title", 23, after=5, line=1.15, color="202020", bold=True)
style("Subtitle", 12, after=10, line=1.15, color="555555")
for name, size, before, after, color in [
    ("Heading 1", 16, 18, 10, "2E74B5"),
    ("Heading 2", 13, 14, 7, "2E74B5"),
    ("Heading 3", 12, 10, 5, "1F4D78"),
]:
    style(name, size, before, after, 1.25, color, True).paragraph_format.keep_with_next = True
for name, size, before, after, line, color, bold in [
    ("Table Text", 9.5, 0, 0, 1.10, "202020", False),
    ("Table Header", 9.5, 0, 0, 1.10, "1F4D78", True),
    ("Metadata", 9.5, 0, 3, 1.15, "555555", False),
    ("Note", 9.5, 4, 4, 1.15, "555555", False),
    ("Table Source", 9, 4, 4, 1.15, "555555", False),
    ("Code Text", 9, 0, 4, 1.15, "202020", False),
    ("Header", 9, 0, 0, 1.10, "666666", False),
    ("Footer", 9, 0, 0, 1.10, "666666", False),
    ("Action List", 11, 0, 4, 1.25, "202020", False),
    ("Lead", 11, 4, 8, 1.25, "9B1C1C", True),
]:
    style(name, size, before, after, line, color, bold)

# Native decimal numbering, with explicit compact-reference indent tokens.
num_root = doc.part.numbering_part.element
abstract_id = (
    max([int(x.get(qn("w:abstractNumId"))) for x in num_root.findall(qn("w:abstractNum"))] + [0])
    + 1
)
num_id = max([int(x.get(qn("w:numId"))) for x in num_root.findall(qn("w:num"))] + [0]) + 1
abstract = OxmlElement("w:abstractNum")
abstract.set(qn("w:abstractNumId"), str(abstract_id))
level = OxmlElement("w:lvl")
level.set(qn("w:ilvl"), "0")
for tag, val in [("start", "1"), ("numFmt", "decimal"), ("lvlText", "%1."), ("lvlJc", "left")]:
    e = OxmlElement("w:" + tag)
    e.set(qn("w:val"), val)
    level.append(e)
ppr = OxmlElement("w:pPr")
ind = OxmlElement("w:ind")
ind.set(qn("w:left"), "540")
ind.set(qn("w:hanging"), "271")
ppr.append(ind)
tabs = OxmlElement("w:tabs")
tab = OxmlElement("w:tab")
tab.set(qn("w:val"), "num")
tab.set(qn("w:pos"), "540")
tabs.append(tab)
ppr.append(tabs)
spacing = OxmlElement("w:spacing")
spacing.set(qn("w:before"), "0")
spacing.set(qn("w:after"), "80")
spacing.set(qn("w:line"), "340")
spacing.set(qn("w:lineRule"), "exact")
ppr.append(spacing)
level.append(ppr)
abstract.append(level)
num_root.append(abstract)
number = OxmlElement("w:num")
number.set(qn("w:numId"), str(num_id))
aid = OxmlElement("w:abstractNumId")
aid.set(qn("w:val"), str(abstract_id))
number.append(aid)
num_root.append(number)


def setup_section(section, landscape=False):
    section.orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
    section.page_width = Inches(11 if landscape else 8.5)
    section.page_height = Inches(8.5 if landscape else 11)
    section.top_margin = section.bottom_margin = Inches(1)
    section.left_margin = section.right_margin = Inches(1)
    section.header_distance = section.footer_distance = Inches(0.492)
    grid = section._sectPr.find(qn("w:docGrid"))
    if grid is not None:
        section._sectPr.remove(grid)


setup_section(doc.sections[0])
hp = doc.sections[0].header.paragraphs[0]
hp.style = doc.styles["Header"]
hp.text = "江苏润盛  |  点位与采集链路验证  |  2026-09-07"
fp = doc.sections[0].footer.paragraphs[0]
fp.style = doc.styles["Footer"]
fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
fp.add_run("项目测试记录  ·  第 ")
for field, suffix in [("PAGE", " 页 / 共 "), ("NUMPAGES", " 页")]:
    el = OxmlElement("w:fldSimple")
    el.set(qn("w:instr"), field)
    run = OxmlElement("w:r")
    tx = OxmlElement("w:t")
    tx.text = "1"
    run.append(tx)
    el.append(run)
    fp._p.append(el)
    fp.add_run(suffix)


def p(text, sty="Normal"):
    return doc.add_paragraph(text, sty)


def h(text, level=1, new_page=False):
    para = doc.add_heading(text, level)
    if new_page:
        para.paragraph_format.page_break_before = True
    return para


def action(text):
    para = p(text, "Action List")
    pr = para._p.get_or_add_pPr()
    np = OxmlElement("w:numPr")
    il = OxmlElement("w:ilvl")
    il.set(qn("w:val"), "0")
    ni = OxmlElement("w:numId")
    ni.set(qn("w:val"), str(num_id))
    np.append(il)
    np.append(ni)
    pr.append(np)


def table(headers, rows, widths, center=()):
    t = doc.add_table(rows=1, cols=len(headers))
    for j, text in enumerate(headers):
        t.rows[0].cells[j].text = str(text)
    for values in rows:
        for j, text in enumerate(values):
            if j == 0:
                cells = t.add_row().cells
            cells[j].text = str(text)
    apply_table_geometry(
        t,
        widths,
        table_width_dxa=sum(widths),
        indent_dxa=120,
        cell_margins_dxa={"top": 80, "bottom": 80, "start": 120, "end": 120},
    )
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement("w:" + side)
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), "4")
        e.set(qn("w:color"), "C8D1DB")
        borders.append(e)
    t._tbl.tblPr.append(borders)
    for i, row in enumerate(t.rows):
        pr = row._tr.get_or_add_trPr()
        pr.append(OxmlElement("w:cantSplit"))
        if i == 0:
            pr.append(OxmlElement("w:tblHeader"))
        for j, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            sh = OxmlElement("w:shd")
            sh.set(qn("w:fill"), "E8EEF5" if i == 0 else "FFFFFF")
            cell._tc.get_or_add_tcPr().append(sh)
            for para in cell.paragraphs:
                para.style = doc.styles["Table Header" if i == 0 else "Table Text"]
                para.paragraph_format.keep_with_next = i == 0
                if i == 0 or j in center:
                    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return t


# Front matter and executive findings, without a separate decorative cover.
p("江苏润盛\n点位表与远程测试报告", "Title")
p("签名部署、实机只读采集、数据库写入复验及未决问题", "Subtitle")
p("记录日期：2026年9月7日    文档版本：1.0", "Metadata")
p("目标机：WIN-OAUCM8UQUGH / 100.109.90.21    软件版本：deploy-20260907.1", "Metadata")
p("证据截止：2026-09-07 11:40:34（北京时间，UTC+8）", "Metadata")
p("本文件归档既有测试结果；编制文档期间未新增远程操作。", "Metadata")
h("一、结论概览")
p(
    "已完成签名部署和本轮批准的验证范围。软件修复与原始值存储通过，但尚未完成全部生产点位验收。",
    "Lead",
)
p(
    "目标机五个服务运行，前端HTTP 200，API与采集网关内部就绪检查通过。读取的36个请求地址原值，已在目标机独立测试数据库完成实时表和历史表的准确写入及回读。"
)
table(
    ["验证层次", "覆盖量", "结论"],
    [
        ["实机只读通信", "3次FC3请求 / 36个请求地址", "CRC有效；零重试；零设备寄存器写入"],
        ["独立测试库", "原值各36条；测试总计126/168条", "实时/历史回读一致"],
        ["候选配置测试", "42项寄存器模拟；4项冲突开关量", "模拟存储与拒绝保护符合预期"],
        ["证据一致性", "57项核验", "全部通过"],
        ["正式生产采集", "账号、设备、点位均未配置", "未启用连续轮询；未完成投产验收"],
    ],
    [2300, 2600, 4460],
)
p("表中126/168分别为测试库实时/历史记录总量，不是客户遥测。", "Table Source")
h("重要边界", 2)
p(
    "46行点位来自旧软件的BCMM、CBMM两种型号候选，不代表已安装46个物理点。设备型号、请求地址含义、符号和工程倍率仍需确认。旧换算出现1000℃、9.9Hz等疑点，不能直接作为有效测量值。"
)
p(
    "阅读顺序：部署与通信 → 数据库及修复 → 未决问题 → 全量候选点位表 → 证据索引。点位表采用横向页面，保留原名称、倍率、单位及逐项结果。",
    "Note",
)

h("二、部署与实机通信", new_page=True)
h("部署结果", 2)
p(
    "新版本于2026-09-07 11:28:19生效，受控升级状态为committed。沿用既有Ed25519签名密钥和发布身份，未绕过签名、完整包文件及镜像身份验证。旧版本deploy-20260905.2及数据库、角色、环境备份均保留。"
)
p(
    "正常升级流程重建五个服务容器；没有重启操作系统或Docker Desktop，没有compose down、卷删除或生产数据删除。环境只有五个镜像引用变化，平台及非发布字段未变。数据库迁移头保持0012_alarm_notification_runtime。"
)
p(
    "发布使用独立干净工作树，未带入其他未提交修改；本轮未推送、创建PR或合并。完整包543,858,089字节，经现有差量方案传输111,421,417字节，重构哈希一致后才执行正常Apply。临时签名进程已停止。",
    "Note",
)
h("物理读取条件", 2)
p(
    "设备通过FTDI USB转RS485适配器连接。USB身份0403:6001/AI06JYFW，串口别名/dev/ruisheng-rs485；9600波特、8数据位、无校验、1停止位，从站地址1，仅使用FC3读取保持寄存器。"
)
p(
    "固定顺序：0..5 → 6..26 → 27..35，后段必须在前段有效后执行；400ms超时、500ms间隔策略，每段最多1次重试，总计最多6次请求，接收上限64字节。原B-06白名单未改。"
)
responses = [r for r in audit if r["event"] == "response_rx"]
requests = [r for r in audit if r["event"] == "request_tx"]
table(
    ["请求范围", "数量", "发送帧（十六进制）", "延迟 / 结果"],
    [
        [
            f"{r['start_address']}..{r['start_address'] + r['register_count'] - 1}",
            r["register_count"],
            r["tx_hex"],
            f"{s['latency_ms']}ms；首次成功",
        ]
        for r, s in zip(requests, responses, strict=True)
    ],
    [1400, 800, 3650, 3510],
    center=(0, 1),
)
p(
    "执行时间11:31:08至11:31:10；3次只读请求均CRC有效，0重试、0设备寄存器写操作。36个原值见后文逐项表。",
    "Table Source",
)
p(
    "审计attempted_write_bytes=24指发送读取命令的串口字节数，不是写设备寄存器。安装工具时外层命令曾返回1；随后通过安装收据、哈希、独立零I/O预演及实际执行返回0确认安装状态，未把外层异常伪报为成功。B-04生产采集门禁仍为BLOCKED。",
    "Note",
)

h("三、数据库验证、修复与回归")
p(
    "2026-09-07 11:36:02至11:36:05，在目标机真实PostgreSQL/Timescale服务上建立唯一独立测试库，使用当前已部署GW镜像执行。未替换内存模块，未修改生产容器文件或生产设备/点位配置。"
)
p("保留测试库：test_point_pipeline_20260907_d135806d", "Code Text")
p(
    "验证链路：Registry → FrameIngestor → BatchWriter → Repository → PostgreSQL/TimescaleDB → SELECT回读。物理原帧回放只保留无符号原值，模拟测试另外验证候选地址和旧倍率的计算过程。"
)
table(
    ["检查项目", "实际结果"],
    [
        ["36个物理原值回放", "实时36、历史36；逐地址回读精确一致"],
        ["42项正数模拟值、倍率与地址", "每项实时1、历史2；最新值和时间戳一致"],
        ["42个旧s16类型输入0xffff", "全部拒绝；实时0、历史0，不误写为正65535"],
        ["4个冲突FC1旧配置", "全部拒绝；实时0、历史0、发布0"],
        ["功能码错配、错误CRC", "均拒绝，无数据库记录"],
        ["取队列记录与停机同时完成", "实时1、历史1，已取出记录不丢失"],
        ["停机排空46条待写记录", "实时46、历史46"],
        ["历史唯一键冲突事务回滚", "实时UPSERT同时回滚；原值保持7"],
    ],
    [3930, 5430],
)
p(
    "测试总计实时126、历史168条，已独立再次查询。物理原值各36条，其余为诊断用例。测试发布器仅记录内存事件，不向生产Redis发布；因此不等于Redis投递、API权限、浏览器流程或持续轮询验收。",
    "Table Source",
)
h("已修复并部署的问题", 2)
p(
    "批量写入：修复队列取值与停机信号同时完成时的丢数问题，先保留已取出的记录再排空队列。采集解码：拒绝尚不支持的类型、冲突开关量配置和点位/响应功能码错配，避免静默写入错误值。此次没有启用旧s16负数解码。"
)
p(
    "认证探测工具新增固定三段FC3白名单，并独立核对顺序、前段成功条件、帧内容、重试与总预算；没有开放任意地址扫描或写功能码。"
)
h("回归和最终状态", 2)
p(
    "发布/签名/升级/Modbus/GW组合654通过、2项POSIX专属跳过；API/shared 593通过、8项既有跳过；独立操作端升级套件60通过。测试范围可能重叠，不作独立总数相加。Ruff、mypy通过。57项证据核验全部通过。",
    "Note",
)
p(
    "11:40:34最终快照：五服务运行、重启计数均0；API与GW内部全ready、Web 200；无探测/实验容器残留及维护租约。生产用户、设备、点位、实时、历史记录均0；订阅、原串口配置和信任密钥未变。",
    "Note",
)

h("四、未决问题与处理顺序")
p(
    "以下为后续待办，不是本轮新增批准或已经执行的操作。当前结论维持“原值链路通过，生产点位验收未完成”。",
    "Lead",
)
issues = [
    (
        "U01  型号与地址响应行为未确认（阻断投产）",
        "BCMM与CBMM仍是备选型号。三个不同请求区间的响应均以[3,0]开头；Modbus读响应不回显起始地址，CRC只能证明帧完整，不能证明设备按起点返回了正确寄存器。需通过有界重叠区间测试和原协议/固件资料核实。",
    ),
    (
        "U02  工程倍率、单位及有符号语义未确认（阻断投产）",
        "地址26原值10000按旧倍率会成为1000℃，地址23原值990会成为9.9Hz。不能据此认定真实温度或频率，也不能擅改倍率凑数。旧s16目前被拒绝，需要先确认协议再决定是否实现和验证负值解码。",
    ),
    (
        "U03  四个FC1候选冲突且未实测（待确认）",
        "旧配置同时使用s16及地址/位号组合，存在解释冲突；只完成软件拒绝测试，未向设备发送FC1请求。应先确认实际型号是否有这些点，再按受控只读范围测试。",
    ),
    (
        "U04  正式采集及完整业务链路未验收（未投用）",
        "生产用户、设备、点位尚未配置，连续轮询未启用。独立数据库回放不能替代生产Redis投递、API授权、浏览器业务操作及数据库持续增量验收。",
    ),
    (
        "U05  长时间通信稳定性未验证（待观察）",
        "本轮仅3次读取全部成功。上一阶段曾有一次截短响应并重试恢复，根因尚未确定；不据此断言长期稳定，也不盲目改变波特率或接线。",
    ),
]
for title, body in issues:
    h(title, 2)
    p(body)
h("建议推进顺序", 2)
action(
    "先核实设备型号和地址相关性：制定并批准固定、低频、有界、无写入的重叠区间读取方案，保留原始请求与响应。"
)
action(
    "再确认点名、单位、倍率及符号语义，解决适用的FC1配置；以协议资料和可控实物参考完成对应验证。"
)
action(
    "最后导入已确认点位，启用经过批准的连续采集，逐点验收生产实时/历史增量、时间戳、重连及前后端显示。"
)

section = doc.add_section(WD_SECTION.NEW_PAGE)
setup_section(section, landscape=True)
h("五、全量候选点位表")
p(
    "地址均为十进制、零基请求地址。42项FC3旧类型均为s16；点偏移0、用户倍率1、用户偏移0。旧倍率及单位仅作资料记录，原值不换算为工程值。全部候选禁止直接导入生产。",
    "Note",
)
p(
    "原值R/H：实时/历史各1条回读一致；不同型号共用地址时引用同一份原值证据。模拟R/H：该候选用受支持的无符号“字”类型及旧倍率输入两次，实时保留最新1条、历史2条。42项旧s16输入0xffff的拒绝测试均通过，不代表已经支持负数。",
    "Note",
)
point_widths = [1250, 2100, 850, 1750, 800, 1300, 1300, 3610]
assert sum(point_widths) == 12960
point_headers = [
    "候选ID",
    "旧名称 / 标签",
    "FC/\n地址",
    "旧倍率 / 单位",
    "原值",
    "原值R/H",
    "模拟R/H",
    "状态 / 未决",
]


def point_row(c):
    a = raw[c["address"]]
    t = tested[c["id"]]
    assert a["realtime_pass"] and a["history_pass"] and t["realtime_pass"] and t["history_pass"]
    note = "型号、映射及工程值未定"
    if c["address"] in (0, 6, 27):
        note = "区间首值重复；映射待核"
    if c["id"] == "CBMM-023":
        note = "疑点：旧换算9.9Hz，不可确认"
    if c["id"] == "CBMM-026":
        note = "疑点：旧换算1000℃，不可采信"
    return [
        c["id"],
        c["name"] + " / " + c["label"],
        f"3/{c['address']}",
        f"{c['ratio']:g} / {c['unit'] or '未注明'}",
        a["raw"],
        "1/1 通过",
        "1/2 通过",
        note,
    ]


h("FC3寄存器候选：42行（BCMM 6行、CBMM 36行）", 2)
table(
    point_headers,
    [point_row(c) for c in candidates if c["fc"] == 3],
    point_widths,
    center=(2, 3, 4, 5, 6),
)
p(
    "BCMM与CBMM在0..5地址重合，不能因此认定两个型号同时存在。地址0、6、27对应各区间首值3，三个响应前两项均[3,0]；旧换算1000℃和9.9Hz均未证实。详见U01、U02。",
    "Table Source",
)
h("CBMM开关量候选：4行，物理采集未测试", 2)
coilrows = []
for c in candidates:
    if c["fc"] != 1:
        continue
    assert tested[c["id"]]["rejection_pass"]
    coilrows.append(
        [
            c["id"],
            c["name"] + " / " + c["label"],
            f"1/{c['address']}/{c['bit']}",
            "未测试",
            "拒绝通过；实时0、历史0、发布0",
            "s16及地址/位号冲突；禁止导入",
        ]
    )
table(
    ["候选ID", "旧名称 / 标签", "FC/地址/位", "实机", "软件与数据库", "未决问题"],
    coilrows,
    [1400, 2450, 1250, 1000, 3350, 3510],
    center=(2, 3),
)
p("以上4行只完成旧配置的拒绝测试，未向设备发送FC1请求；不得将其标为实机采集通过。", "Table Source")

section = doc.add_section(WD_SECTION.NEW_PAGE)
setup_section(section)
h("六、证据与追溯信息")
p(
    "证据截止2026-09-07 11:40:34。以下索引指向项目内保留记录，Word文档汇总主要结论和全量候选行；具体帧、逐条模拟值和数据库时间戳以原始JSON/JSONL为准。",
    "Note",
)
p("本地证据目录：", "Metadata")
p("D:/江苏润盛/docs/superpowers/specs/evidence/point-release-20260907", "Code Text")
table(
    ["文件", "内容"],
    [
        ["POINT-TABLE.md", "46行候选映射及逐项结果"],
        ["README.md", "本轮签名部署和复验记录"],
        ["target-pipeline-results.json", "实际部署代码的逐项原值、模拟值、拒绝及数据库回读"],
        ["target-final-state.json", "最终服务、镜像、哈希、生产计数及健康快照"],
        ["validation-summary.json", "57项一致性核验结果及目标机副本哈希"],
        ["release-build.json", "签名发布标识、构建和本地回归信息"],
        ["point-release-20260907-physical.jsonl", "三段请求、完整响应、CRC、原值与时间戳"],
        ["modbus-probe-release.json", "探测工具安装收据，绑定当前镜像"],
    ],
    [4200, 5160],
)
p(
    "同目录另保留full-upgrade操作日志及modbus-runner预演/实测审计，文件名带下列对应操作ID。",
    "Table Source",
)
h("版本与运行标识", 2)
p("源提交：" + release["source_commit"], "Code Text")
p("逻辑标识：" + release["logical_identity"], "Code Text")
p("部署操作：1c95b8c8-8f92-4e8b-962f-207800c7e9c9", "Code Text")
p("实机运行：" + result["physical_run_id"], "Code Text")
p("预演运行：ba1ac2a2-f5da-42b9-b25f-aa3dfe088229", "Code Text")
p("发布身份：ruisheng-release；命名空间：ruisheng-candidate-v1", "Code Text")
p("信任指纹：" + release["key_fingerprint"], "Code Text")
h("目标机保留材料", 2)
p("当前候选：C:/Ruisheng/candidates/deploy-20260907.1", "Code Text")
p("站点：C:/Ruisheng/candidates/site-deploy-20260831.1", "Code Text")
p(
    "备份位于上述站点的backups/部署操作ID目录，含ruisheng.dump、roles.sql、.env.prod.before。原配置和旧版本保留，未将环境或连接密码纳入本文件。",
    "Note",
)
p(
    "原始候选来源：相邻point-pipeline-20260907/candidate-summary.json；更早的旧软件提取证据位于b08-20260905/legacy-point-candidates.json。上一阶段报告保留了旧版失败与临时修复对照，本文件以本轮已部署复验结果为准。",
    "Note",
)

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
assert audit_docx_tables(OUTPUT) == 0

# Audit the generated artifact against structured evidence and preset geometry.
loaded = Document(OUTPUT)
ids = []
for t in loaded.tables:
    for row in t.rows[1:]:
        if row.cells[0].text in tested:
            ids.append(row.cells[0].text)
assert len(ids) == 46 and len(set(ids)) == 46 and set(ids) == set(tested)
for s in loaded.sections:
    assert s.left_margin.twips == 1440 and s.right_margin.twips == 1440
    assert s.top_margin.twips == 1440 and s.bottom_margin.twips == 1440
    assert (s.page_width.twips, s.page_height.twips) in {(12240, 15840), (15840, 12240)}
assert len(loaded.sections) == 3
assert loaded.styles["Normal"].paragraph_format.space_after.pt == 6
assert loaded.styles["Normal"].paragraph_format.line_spacing.pt == 17
summary = {
    "output": str(OUTPUT),
    "candidate_rows": len(ids),
    "unique_candidate_rows": len(set(ids)),
    "physical_requested_addresses": len(raw),
    "tables": len(loaded.tables),
    "sections": len(loaded.sections),
    "table_geometry_issues": 0,
    "preset": "compact_reference_guide",
    "header": "memo_masthead",
    "source_sha256": {
        path.name: hashlib.sha256(path.read_bytes()).hexdigest()
        for path in [
            EVIDENCE / "README.md",
            EVIDENCE / "POINT-TABLE.md",
            EVIDENCE / "target-pipeline-results.json",
            EVIDENCE / "target-final-state.json",
            EVIDENCE / "validation-summary.json",
            PREVIOUS / "candidate-summary.json",
        ]
    },
}
(QA / "structural-checks.json").write_text(
    json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8"
)
print(json.dumps(summary, ensure_ascii=True))
