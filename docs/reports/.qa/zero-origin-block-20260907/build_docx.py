"""Create the B-11 Word report from retained evidence, without remote I/O.

Preset compact_reference_guide; memo_masthead without rule.
Named overrides: Microsoft YaHei East Asian font, exact Chinese line metrics
(body/list 17 pt, table 13 pt, title 29 pt, headings 21/18/17 pt), table text
9.5 pt, metadata 9.5 pt, evidence text 9.5 pt, header/footer 9 pt.
"""

from __future__ import annotations

import hashlib
import json
import re
import sys
from datetime import UTC, datetime
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path("D:/江苏润盛")
EVIDENCE = ROOT / "docs/superpowers/specs/evidence/zero-origin-block-20260907"
SOURCE = EVIDENCE / "RUN-RESULTS.md"
QA = Path(__file__).parent
OUTPUT = ROOT / "docs/reports/江苏润盛_零起点整块读取测试报告_20260907.docx"
SKILL = Path(
    "C:/Users/admin/.codex/plugins/cache/openai-primary-runtime/documents/26.826.12353/skills/documents"
)
sys.path.insert(0, str(SKILL / "scripts"))
from table_geometry import apply_table_geometry, audit_docx_tables


def read_json(name):
    return json.loads((EVIDENCE / name).read_text(encoding="utf-8-sig"))


comparison = read_json("whole-block-comparison.json")
validation = read_json("completed-run-validation.json")
assert comparison["completed_tx"] == 6 and comparison["retries"] == 1
assert validation["check_count"] == 88 and all(validation["checks"].values())
assert comparison["formal_acceptance"] == "BLOCKED"
doc = Document()
doc.core_properties.title = "江苏润盛：零起点整块读取测试报告"
doc.core_properties.subject = "B-11 实机只读通信、原值对照与未决问题"
doc.core_properties.author = doc.core_properties.last_modified_by = "项目测试记录"
doc.core_properties.created = datetime(2026, 9, 7, tzinfo=UTC)
doc.core_properties.modified = datetime.now(UTC)
doc.core_properties.keywords = "B-11,36寄存器,只读测试,正式验收未通过"


def style(name, size, before=0, after=6, leading=17, color="202020", bold=False):
    s = (
        doc.styles[name]
        if name in doc.styles
        else doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    )
    s.font.name, s.font.size = "Calibri", Pt(size)
    s.font.bold, s.font.italic, s.font.underline = bold, False, False
    s.font.color.rgb = RGBColor.from_string(color)
    fonts = s.element.get_or_add_rPr().get_or_add_rFonts()
    for key in ("ascii", "hAnsi", "cs"):
        fonts.set(qn("w:" + key), "Calibri")
    fonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    for key in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        fonts.attrib.pop(qn("w:" + key), None)
    fmt = s.paragraph_format
    fmt.space_before, fmt.space_after = Pt(before), Pt(after)
    fmt.line_spacing, fmt.widow_control = Pt(leading), True
    fmt.alignment = WD_ALIGN_PARAGRAPH.LEFT
    fmt.keep_with_next = name.startswith("Heading")
    ppr = s.element.get_or_add_pPr()
    for border in list(ppr.findall(qn("w:pBdr"))):
        ppr.remove(border)
    snap = OxmlElement("w:snapToGrid")
    snap.set(qn("w:val"), "0")
    ppr.append(snap)
    return s


style("Normal", 11)
style("Title", 23, after=5, leading=29, bold=True)
style("Subtitle", 12, after=10, leading=16, color="555555")
style("Heading 1", 16, 18, 10, 21, "2E74B5", True)
style("Heading 2", 13, 14, 7, 18, "2E74B5", True)
style("Heading 3", 12, 10, 5, 17, "1F4D78", True)
style("Table Text", 9.5, after=0, leading=13)
style("Table Header", 9.5, after=0, leading=13, color="1F4D78", bold=True)
style("Metadata", 9.5, after=3, leading=13.5, color="555555")
style("Table Source", 9, 4, 4, 13, "555555")
style("Evidence Text", 9.5, after=4, leading=14, color="555555")
style("Action List", 11, after=4)
style("Lead", 11, 4, 8, 17, "9B1C1C", True)
style("Header", 9, after=0, leading=12, color="666666")
style("Footer", 9, after=0, leading=12, color="666666")
section = doc.sections[0]
section.page_width, section.page_height = Inches(8.5), Inches(11)
section.top_margin = section.bottom_margin = Inches(1)
section.left_margin = section.right_margin = Inches(1)
section.header_distance = section.footer_distance = Inches(0.492)
grid = section._sectPr.find(qn("w:docGrid"))
if grid is not None:
    section._sectPr.remove(grid)
header = section.header.paragraphs[0]
header.style = doc.styles["Header"]
header.text = "江苏润盛  |  B-11 零起点整块读取测试  |  2026-09-07"
footer = section.footer.paragraphs[0]
footer.style, footer.alignment = doc.styles["Footer"], WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run("测试记录  |  第 ")
for field, suffix in (("PAGE", " 页 / 共 "), ("NUMPAGES", " 页")):
    el = OxmlElement("w:fldSimple")
    el.set(qn("w:instr"), field)
    run, text = OxmlElement("w:r"), OxmlElement("w:t")
    text.text = "1"
    run.append(text)
    el.append(run)
    footer._p.append(el)
    footer.add_run(suffix)
update = OxmlElement("w:updateFields")
update.set(qn("w:val"), "true")
doc.settings.element.append(update)


def numbering(fmt):
    root = doc.part.numbering_part.element
    aid = max(int(x.get(qn("w:abstractNumId"))) for x in root.findall(qn("w:abstractNum"))) + 1
    nid = max(int(x.get(qn("w:numId"))) for x in root.findall(qn("w:num"))) + 1
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(aid))
    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    for tag, val in (
        ("start", "1"),
        ("numFmt", fmt),
        ("lvlText", "%1." if fmt == "decimal" else "\u2022"),
        ("lvlJc", "left"),
    ):
        child = OxmlElement("w:" + tag)
        child.set(qn("w:val"), val)
        level.append(child)
    ppr, ind = OxmlElement("w:pPr"), OxmlElement("w:ind")
    ind.set(qn("w:left"), "540")
    ind.set(qn("w:hanging"), "271")
    ppr.append(ind)
    tabs, tab = OxmlElement("w:tabs"), OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "540")
    tabs.append(tab)
    ppr.append(tabs)
    spacing = OxmlElement("w:spacing")
    for key, value in (("before", "0"), ("after", "80"), ("line", "340"), ("lineRule", "exact")):
        spacing.set(qn("w:" + key), value)
    ppr.append(spacing)
    level.append(ppr)
    abstract.append(level)
    root.append(abstract)
    num, aid_node = OxmlElement("w:num"), OxmlElement("w:abstractNumId")
    num.set(qn("w:numId"), str(nid))
    aid_node.set(qn("w:val"), str(aid))
    num.append(aid_node)
    root.append(num)
    return nid


NUMBERS = {"decimal": numbering("decimal"), "bullet": numbering("bullet")}
link_targets, expected_paragraphs, table_data = [], [], []


def inline(paragraph, text):
    for piece in re.split(r"(\[[^\]]+\]\([^)]+\)|\*\*[^*]+\*\*)", text):
        if not piece:
            continue
        match = re.fullmatch(r"\[([^\]]+)\]\(([^)]+)\)", piece)
        if match:
            label, target = match.groups()
            path = (EVIDENCE / target).resolve()
            assert path.is_file(), path
            link_targets.append(str(path))
            rid = paragraph.part.relate_to(path.as_uri(), RT.HYPERLINK, is_external=True)
            link, run, prop, color, node = [
                OxmlElement("w:" + t) for t in ("hyperlink", "r", "rPr", "color", "t")
            ]
            link.set(qn("r:id"), rid)
            color.set(qn("w:val"), "2E74B5")
            prop.append(color)
            run.append(prop)
            node.text = label
            run.append(node)
            link.append(run)
            paragraph._p.append(link)
        elif piece.startswith("**") and piece.endswith("**"):
            paragraph.add_run(piece[2:-2]).bold = True
        else:
            paragraph.add_run(piece)


def para(text, style_name="Normal", list_kind=None):
    p = doc.add_paragraph(style=style_name)
    inline(p, text)
    if list_kind:
        pr, il, num = OxmlElement("w:numPr"), OxmlElement("w:ilvl"), OxmlElement("w:numId")
        il.set(qn("w:val"), "0")
        num.set(qn("w:val"), str(NUMBERS[list_kind]))
        pr.append(il)
        pr.append(num)
        p._p.get_or_add_pPr().append(pr)
    expected_paragraphs.append(re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text).replace("**", ""))
    return p


def add_table(lines):
    matrix = [[value.strip() for value in line.strip().strip("|").split("|")] for line in lines]
    assert all(re.fullmatch(r":?-+:?", col) for col in matrix[1])
    headers, rows = matrix[0], matrix[2:]
    assert all(len(row) == len(headers) == 5 for row in rows)
    ti = len(table_data)
    widths = [
        [700, 1250, 1250, 2050, 4110],
        [3560, 1450, 1450, 1450, 1450],
        [1200, 1350, 1350, 1350, 4110],
    ][ti]
    if ti == 1:
        headers = ["对照范围", "相同非零", "相同零值", "动态未决", "稳定不一致"]
    table = doc.add_table(rows=1, cols=5)
    for j, value in enumerate(headers):
        table.rows[0].cells[j].text = value
    for row in rows:
        for cell, value in zip(table.add_row().cells, row, strict=True):
            cell.text = value
    apply_table_geometry(
        table,
        widths,
        table_width_dxa=9360,
        indent_dxa=120,
        cell_margins_dxa={"top": 80, "bottom": 80, "start": 120, "end": 120},
    )
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement("w:" + side)
        for key, value in (("val", "single"), ("sz", "4"), ("color", "C8D1DB")):
            e.set(qn("w:" + key), value)
        borders.append(e)
    table._tbl.tblPr.append(borders)
    for index, row in enumerate(table.rows):
        pr = row._tr.get_or_add_trPr()
        pr.append(OxmlElement("w:cantSplit"))
        if index == 0:
            pr.append(OxmlElement("w:tblHeader"))
        for column, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            shd = OxmlElement("w:shd")
            shd.set(qn("w:fill"), "E8EEF5" if index == 0 else "FFFFFF")
            cell._tc.get_or_add_tcPr().append(shd)
            for p in cell.paragraphs:
                p.style = doc.styles["Table Header" if index == 0 else "Table Text"]
                p.paragraph_format.keep_with_next = index == 0
                if index == 0 or (ti == 1 and column > 0) or (ti != 1 and column < 4):
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    table_data.append(rows)
    if ti == 1:
        para(
            "相同非零：前、中、后均相同的非零原值；相同零值不证明点位身份。动态未决表示前后参考已变化。",
            "Table Source",
        )
    elif ti == 2:
        para(
            "数据源：physical.jsonl 与 whole-block-comparison.json。表内均为原始字，不是已确认的工程测量值。",
            "Table Source",
        )


para("江苏润盛\n零起点整块读取测试报告", "Title")
para("B-11 实机只读通信、36 个响应位置原值对照及未决问题", "Subtitle")
para("测试日期：2026年9月7日    文档版本：1.0", "Metadata")
para("目标机：WIN-OAUCM8UQUGH / 100.109.90.21", "Metadata")
para("证据截止：2026-09-07 13:31:18（北京时间，UTC+8）", "Metadata")
para("本文件整理既有测试证据；编制文档期间未新增远程测试或设备操作。", "Metadata")

# The retained source has only headings, prose, flat lists and pipe tables.
lines = SOURCE.read_text(encoding="utf-8-sig").splitlines()
index, current_heading = 0, ""
while index < len(lines):
    line = lines[index].strip()
    index += 1
    if not line or line.startswith("# "):
        continue
    if line.startswith("## "):
        current_heading = line[3:]
        heading = doc.add_heading(current_heading, 1)
        continue
    if line.startswith("|"):
        block = [line]
        while index < len(lines) and lines[index].strip().startswith("|"):
            block.append(lines[index].strip())
            index += 1
        add_table(block)
        continue
    if line.startswith("- "):
        sty = "Evidence Text" if current_heading == "证据入口" else "Action List"
        para(line[2:], sty, "bullet")
    elif re.match(r"^\d+\. ", line):
        para(re.sub(r"^\d+\. ", "", line), "Action List", "decimal")
    else:
        assert not line.startswith((chr(96) * 3, ">", "###", "![")), line
        para(line, "Lead" if line.startswith("**") else "Normal")

assert len(table_data) == 3 and len(table_data[2]) == 36
for row, observation in zip(table_data[2], comparison["full_block_repeatability"], strict=True):
    assert int(row[0]) == observation["response_position"]
    assert [int(value) for value in row[1:4]] == observation["raw_values"]
assert [[int(x) for x in row[1:]] for row in table_data[1]] == [
    [5, 12, 10, 0],
    [1, 2, 3, 0],
    [6, 14, 13, 0],
]
OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
loaded = Document(OUTPUT)
assert [[c.text for c in row.cells] for row in loaded.tables[2].rows[1:]] == table_data[2]
all_text = "\n".join(p.text for p in loaded.paragraphs)
assert all(text in all_text for text in expected_paragraphs)
assert "正式采集验收仍然阻断" in all_text and "537项通过" in all_text
checks = {
    "output": str(OUTPUT),
    "source": str(SOURCE),
    "source_sha256": hashlib.sha256(SOURCE.read_bytes()).hexdigest(),
    "docx_sha256": hashlib.sha256(OUTPUT.read_bytes()).hexdigest(),
    "preset": "compact_reference_guide",
    "header": "memo_masthead",
    "table_count": len(loaded.tables),
    "raw_rows": 36,
    "all_source_paragraphs_preserved": True,
    "raw_values_equal_json_evidence": True,
    "formal_acceptance": "BLOCKED",
    "evidence_links_existing": len(link_targets),
    "table_geometry": audit_docx_tables(OUTPUT),
}
(QA / "structural-checks.json").write_text(
    json.dumps(checks, ensure_ascii=False, indent=2), encoding="utf-8"
)
print(json.dumps(checks, ensure_ascii=False, indent=2))
